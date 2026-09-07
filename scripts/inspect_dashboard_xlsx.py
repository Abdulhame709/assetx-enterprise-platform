import json
import sys
from openpyxl import load_workbook

if len(sys.argv) != 2:
    raise SystemExit("Usage: python3 inspect_dashboard_xlsx.py <file.xlsx>")

workbook = load_workbook(sys.argv[1], read_only=True, data_only=True)
sheet = workbook.active
rows = list(sheet.iter_rows(values_only=True))
print(json.dumps({
    "sheet": sheet.title,
    "firstRows": [["" if value is None else str(value) for value in row] for row in rows[:8]],
    "rowCount": len(rows),
}, ensure_ascii=False))
